#!/usr/bin/env python3
"""Generate the IMF PN-vs-SR AI-comparison slide deck (16:9 PPTX)."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
INK      = RGBColor(0x0B, 0x0B, 0x0B)
SECOND   = RGBColor(0x52, 0x51, 0x4E)
MUTED    = RGBColor(0x89, 0x87, 0x81)
HAIRLINE = RGBColor(0xE1, 0xE0, 0xD9)
SURFACE  = RGBColor(0xFC, 0xFC, 0xFB)
CARD     = RGBColor(0xFF, 0xFF, 0xFF)
BAND     = RGBColor(0xF4, 0xF4, 0xF1)

BLUE     = RGBColor(0x2A, 0x78, 0xD6)
BLUE_D   = RGBColor(0x10, 0x42, 0x81)   # step 650
BLUE_DD  = RGBColor(0x0D, 0x36, 0x6B)   # step 700
AQUA     = RGBColor(0x1B, 0xAF, 0x7A)
AQUA_D   = RGBColor(0x0E, 0x6B, 0x4A)
YELLOW   = RGBColor(0xED, 0xA1, 0x00)
YELLOW_D = RGBColor(0x8A, 0x5D, 0x00)
RED      = RGBColor(0xE3, 0x49, 0x48)
RED_D    = RGBColor(0x9E, 0x2A, 0x29)
VIOLET   = RGBColor(0x4A, 0x3A, 0xA7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

# ordinal blue ramp for pipeline stages (steps 250..700)
RAMP = [RGBColor(0x86, 0xB6, 0xEF), RGBColor(0x55, 0x98, 0xE7),
        RGBColor(0x2A, 0x78, 0xD6), RGBColor(0x1C, 0x5C, 0xAB),
        RGBColor(0x10, 0x42, 0x81), RGBColor(0x0D, 0x36, 0x6B)]

FONT = "Segoe UI"
SW, SH = Inches(13.333), Inches(7.5)
ML = Inches(0.62)                      # left margin
CW = Inches(13.333 - 1.24)             # content width

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def slide_new():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = SURFACE
    return s

def _fmt(r, size, color, bold=False, italic=False, name=FONT):
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.bold = bold; r.font.italic = italic; r.font.name = name

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf

def para(tf, text, size=14, color=SECOND, bold=False, italic=False,
         align=PP_ALIGN.LEFT, space_after=4, space_before=0, first=False,
         line=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    if isinstance(text, str):
        text = [(text, {})]
    for t, kw in text:
        r = p.add_run(); r.text = t
        _fmt(r, kw.get("size", size), kw.get("color", color),
             kw.get("bold", bold), kw.get("italic", italic))
    return p

def header(s, kicker, title, title_color=INK):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, ML, Inches(0.52), Inches(0.32), Inches(0.09))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background(); bar.shadow.inherit = False
    _, tf = tb(s, ML + Inches(0.45), Inches(0.42), CW, Inches(0.3))
    para(tf, kicker.upper(), size=11, color=BLUE, bold=True, first=True)
    _, tf = tb(s, ML, Inches(0.78), CW, Inches(0.62))
    para(tf, title, size=26, color=title_color, bold=True, first=True)

def footer(s):
    n = len(prs.slides._sldIdLst)
    _, tf = tb(s, ML, Inches(7.12), Inches(8), Inches(0.3))
    para(tf, "AI document comparison · IMF Policy Note vs Staff Report", size=9, color=MUTED, first=True)
    _, tf = tb(s, SW - Inches(1.1), Inches(7.12), Inches(0.5), Inches(0.3))
    para(tf, str(n), size=9, color=MUTED, align=PP_ALIGN.RIGHT, first=True)

def box(s, x, y, w, h, fill=CARD, line_c=HAIRLINE, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_c is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_c; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = tf.margin_bottom = Inches(0.09)
    return sp, tf

def card(s, x, y, w, h, accent, title, lines, title_size=14, body_size=12.5,
         title_color=INK, body_color=SECOND):
    """White card with a colored left accent bar, a bold title and body lines."""
    sp, tf = box(s, x, y, w, h)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.12), Inches(0.055), h - Inches(0.24))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent
    bar.line.fill.background(); bar.shadow.inherit = False
    tf.margin_left = Inches(0.22)
    para(tf, title, size=title_size, color=title_color, bold=True, first=True, space_after=3)
    for ln in lines:
        para(tf, ln, size=body_size, color=body_color, space_after=2, line=1.06)
    return sp

def strip(s, y, text_runs, fill=BLUE_DD, h=Inches(0.62), text_color=WHITE, size=13.5):
    sp, tf = box(s, ML, y, CW, h, fill=fill, line_c=None)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text_runs, size=size, color=text_color, first=True, align=PP_ALIGN.LEFT)
    return sp

def arrow(s, x1, y1, x2, y2, color=MUTED, w=1.5, dash=None):
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color; conn.line.width = Pt(w)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    if dash:
        d = ln.makeelement(qn('a:prstDash'), {'val': dash}); ln.insert(0, d)
    return conn

def chip(s, x, y, w, text, fill, text_color=WHITE, size=10.5, h=Inches(0.3)):
    sp, tf = box(s, x, y, w, h, fill=fill, line_c=None, radius=0.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = 0
    para(tf, text, size=size, color=text_color, bold=True, first=True, align=PP_ALIGN.CENTER)
    return sp

# ================================================================ SLIDE 1
s = slide_new()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.14))
band.fill.solid(); band.fill.fore_color.rgb = BLUE_DD; band.line.fill.background(); band.shadow.inherit = False

_, tf = tb(s, ML, Inches(1.5), CW, Inches(0.35))
para(tf, "AI FOR DOCUMENT REVIEW · A CASE STUDY", size=13, color=BLUE, bold=True, first=True)
_, tf = tb(s, ML, Inches(1.95), CW, Inches(1.6))
para(tf, "One Prompt vs. a Pipeline", size=44, color=INK, bold=True, first=True)
_, tf = tb(s, ML, Inches(3.0), Inches(9.2), Inches(1.0))
para(tf, "Using AI to find every material change between an IMF Policy Note and its Staff Report — fiscal and debt policy focus",
     size=18, color=SECOND, first=True, line=1.15)

# small visual: two docs and a delta
dy = Inches(4.55); dw, dh = Inches(1.7), Inches(1.15)
d1, tf1 = box(s, ML, dy, dw, dh, fill=CARD, line_c=BLUE, line_w=1.5)
tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf1, "Policy Note", size=13, color=INK, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=1)
para(tf1, "~80 pp · pre-mission", size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
d2, tf2 = box(s, ML + Inches(3.0), dy, dw, dh, fill=CARD, line_c=BLUE_D, line_w=1.5)
tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf2, "Staff Report", size=13, color=INK, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=1)
para(tf2, "~80 pp · post-negotiation", size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
arrow(s, ML + dw + Inches(0.12), dy + Inches(0.58), ML + Inches(2.88), dy + Inches(0.58), color=SECOND, w=1.75)
_, tf = tb(s, ML + Inches(5.0), dy + Inches(0.18), Inches(6.6), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Goal: ", {"bold": True, "color": INK}),
          ("identify ALL changes with implications for the economic assessment, policy advice, or program design & conditionality", {})],
     size=14, color=SECOND, first=True, line=1.12)

_, tf = tb(s, ML, Inches(6.6), CW, Inches(0.3))
para(tf, "July 2026", size=11, color=MUTED, first=True)

# ================================================================ SLIDE 2
s = slide_new()
header(s, "The task", "Compare two dense IMF documents — and miss nothing")

cy = Inches(1.7); ch = Inches(1.95); cw2 = Inches(3.75)
card(s, ML, cy, cw2, ch, BLUE, "Policy Note", [
    "Staff's intended strategy, written pre-mission",
    "~80 pages",
    "Complex fiscal, program and DSA tables"])
card(s, ML, cy + ch + Inches(0.25), cw2, ch, BLUE_D, "Staff Report", [
    "The post-negotiation outcome, presented to the Board",
    "Different structure: Staff Appraisal, MEFP/LOI, program tables",
    "Not simply “version 2” of the Policy Note"])

rx = ML + cw2 + Inches(0.35); rw = CW - cw2 - Inches(0.35)
sp, tf = box(s, rx, cy, rw, Inches(4.15))
para(tf, "What “ALL changes” means here", size=15, color=INK, bold=True, first=True, space_after=8)
for t in ["Deviations, additions and omissions — in text, tables, footnotes and annexes",
          "Shifts in tone, strength, intensity and emphasis of the discussion",
          "Anything that could affect the economic assessment, policy advice, or program design and conditionality (fiscal & debt)"]:
    para(tf, [("▪  ", {"color": BLUE, "bold": True}), (t, {})], size=13.5, color=SECOND, space_after=10, line=1.15)
para(tf, "Read as an expert IMF economist would — not as a text diff.", size=12.5, color=MUTED, italic=True, space_before=6)

strip(s, Inches(6.15), [("“ALL” makes this a ", {}),
                        ("recall problem", {"bold": True}),
                        (", not a summarization problem — and exhaustive recall is exactly where LLMs are weakest.", {})])
footer(s)

# ================================================================ SLIDE 3
s = slide_new()
header(s, "The challenge", "Long documents play against how LLMs read")

# --- attention curve chart (left ~55%)
px, py = ML, Inches(1.75)
pw, ph = Inches(6.6), Inches(4.0)
plot, _ = box(s, px, py, pw, ph, fill=CARD, line_c=HAIRLINE)
ix, iy = px + Inches(0.55), py + Inches(0.55)      # inner plot origin
iw, ih = pw - Inches(1.0), ph - Inches(1.25)
# gridlines
for frac in (0.0, 0.5, 1.0):
    g = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ix, iy + int(ih * frac), ix + iw, iy + int(ih * frac))
    g.line.color.rgb = HAIRLINE; g.line.width = Pt(0.75); g.shadow.inherit = False
# middle danger band
bx = ix + int(iw * 0.30)
bw = int(iw * 0.40)
bandm = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, iy, bw, ih)
bandm.fill.solid(); bandm.fill.fore_color.rgb = RGBColor(0xFB, 0xE8, 0xE8)
bandm.line.fill.background(); bandm.shadow.inherit = False
# curve: attention high at start, lowest in middle, moderately high at end
pts = []
N = 60
for i in range(N + 1):
    t = i / N
    att = 0.12 + 0.88 * ((2 * t - 1) ** 2) * (0.55 + 0.45 * (1 - t))
    x = int(ix + iw * t)
    y = int(iy + ih * (1 - att))
    pts.append((x, y))
fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1)
fb.add_line_segments(pts[1:], close=False)
curve = fb.convert_to_shape()
curve.fill.background()
curve.line.color.rgb = BLUE; curve.line.width = Pt(2.5)
curve.shadow.inherit = False
# axes
ax = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ix, iy + ih, ix + iw, iy + ih)
ax.line.color.rgb = MUTED; ax.line.width = Pt(1.0); ax.shadow.inherit = False
ay = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ix, iy, ix, iy + ih)
ay.line.color.rgb = MUTED; ay.line.width = Pt(1.0); ay.shadow.inherit = False
# labels
_, tf = tb(s, ix - Inches(0.5), iy - Inches(0.42), Inches(4), Inches(0.3))
para(tf, "Model attention", size=11, color=MUTED, first=True)
_, tf = tb(s, ix, iy + ih + Inches(0.08), iw, Inches(0.3))
para(tf, "Position in an 80-page document", size=11, color=MUTED, first=True, align=PP_ALIGN.CENTER)
_, tf = tb(s, ix + Inches(0.22), iy + Inches(0.02), Inches(2.5), Inches(0.3))
para(tf, "beginning: read closely", size=10.5, color=BLUE_D, bold=True, first=True)
_, tf = tb(s, ix + iw - Inches(1.75), iy + int(ih * 0.24), Inches(1.7), Inches(0.3))
para(tf, "end: read", size=10.5, color=BLUE_D, bold=True, first=True, align=PP_ALIGN.RIGHT)
_, tf = tb(s, bx, iy + Inches(0.28), bw, Inches(0.6))
para(tf, "middle: skimmed", size=10.5, color=RED_D, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=1)
para(tf, "fiscal tables · annexes · DSA live here", size=9.5, color=RED_D, align=PP_ALIGN.CENTER)

# --- right bullets
rx = px + pw + Inches(0.35); rw = SW - rx - ML
card(s, rx, Inches(1.75), rw, Inches(1.22), RED,
     "Changes hide in the worst places",
     ["Tables, footnotes and annexes — the content LLMs handle least reliably"])
card(s, rx, Inches(3.12), rw, Inches(1.22), YELLOW_D,
     "Tone shifts are subtle",
     ["“Should” → “could” matters for policy; no mechanical diff detects it"])
card(s, rx, Inches(4.49), rw, Inches(1.26), BLUE,
     "Fluent ≠ exhaustive",
     ["LLMs are built to answer convincingly; completeness must be engineered"])

_, tf = tb(s, ML, Inches(5.98), CW, Inches(0.6))
para(tf, "Microsoft's own guidance: on long files Copilot may focus on the beginning and ignore later content — split long documents and review them in parts.",
     size=11.5, color=MUTED, italic=True, first=True)
footer(s)

# ================================================================ SLIDE 4
s = slide_new()
header(s, "Under the hood", "Copilot doesn’t “read” your PDF — it searches it")

fy = Inches(2.0); fh = Inches(1.3)
xs = [ML + Inches(2.48) * i for i in range(5)]
bw4 = Inches(2.16)
labels = [("80-page PDF", "attached to the chat", BLUE),
          ("Split into chunks", "fragments, not the whole file", BLUE),
          ("Search index", "chunks stored for retrieval", BLUE_D),
          ("Question retrieves a few chunks", "only the “most relevant” return", BLUE_D),
          ("LLM answers from fragments", "segment-based, not end-to-end", BLUE_DD)]
for i, (t, sub, col) in enumerate(labels):
    shape_kind = MSO_SHAPE.CAN if i == 2 else MSO_SHAPE.ROUNDED_RECTANGLE
    sp, tf = box(s, xs[i], fy, bw4, fh, fill=col, line_c=None, shape=shape_kind)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, t, size=12.5, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=2)
    para(tf, sub, size=10, color=RGBColor(0xDD, 0xE9, 0xFA), align=PP_ALIGN.CENTER)
    if i > 0:
        arrow(s, xs[i] - Inches(0.28), fy + Inches(0.65), xs[i] - Inches(0.04), fy + Inches(0.65), color=SECOND, w=1.75)

# chunk visual under box 2
for k in range(3):
    csp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xs[1] + Inches(0.35 + 0.55 * k), fy + fh + Inches(0.18), Inches(0.45), Inches(0.22))
    csp.fill.solid(); csp.fill.fore_color.rgb = RGBColor(0x9E, 0xC5, 0xF4)
    csp.line.fill.background(); csp.shadow.inherit = False

# advertised vs observed
cy4 = Inches(4.05); ch4 = Inches(1.35)
card(s, ML, cy4, Inches(5.9), ch4, AQUA_D, "Advertised",
     ["Up to ~300 pages / 1.5 million words per file"])
card(s, ML + Inches(6.15), cy4, Inches(5.9), ch4, RED, "Observed in practice",
     ["Text can silently drop from ~page 20–45, depending on density and app",
      "Limits are undocumented and drift over time"])

strip(s, Inches(5.75), [("No end-to-end reading ", {"bold": True}),
                        ("→ a footnote that contradicts a claim 60 pages later is structurally invisible to the model.", {})])
footer(s)

# ================================================================ SLIDE 4b — chat conversation memory
s = slide_new()
header(s, "Under the hood · 2", "The obvious workaround fails: the chat window forgets")

_, tf = tb(s, ML, Inches(1.56), CW, Inches(0.35))
para(tf, "The tempting fix for truncation — split the PDFs yourself and feed the pieces through one chat, then ask:",
     size=13, color=SECOND, first=True)

# chunk row: 8 chunks fed one turn at a time; only the last 5 stay in memory
ry = Inches(2.32); rh5 = Inches(0.6)
cwid = Inches(0.92); pitch5 = Inches(1.12)
N_CH = 8; KEEP_FROM = 3          # chunks 0-2 evicted, 3-7 retained
for i in range(N_CH):
    x = ML + pitch5 * i
    if i < KEEP_FROM:
        sp, tf = box(s, x, ry, cwid, rh5, fill=BAND, line_c=MUTED, line_w=1.0)
        ln = sp.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
        tcol = MUTED
    else:
        sp, tf = box(s, x, ry, cwid, rh5, fill=BLUE, line_c=None)
        tcol = WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.02)
    para(tf, f"chunk {i+1}", size=10.5, color=tcol, bold=True, first=True, align=PP_ALIGN.CENTER)

# memory window bracket around retained chunks
wx = ML + pitch5 * KEEP_FROM - Inches(0.09)
ww = pitch5 * (N_CH - 1 - KEEP_FROM) + cwid + Inches(0.18)
win, _ = box(s, wx, ry - Inches(0.13), ww, rh5 + Inches(0.26), fill=None, line_c=BLUE_D, line_w=1.75, radius=0.12)
ln = win.line._get_or_add_ln()
ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
_, tf = tb(s, wx, ry - Inches(0.48), ww, Inches(0.3))
para(tf, "conversation memory — fixed size", size=10.5, color=BLUE_D, bold=True, first=True, align=PP_ALIGN.CENTER)
_, tf = tb(s, ML, ry + rh5 + Inches(0.18), pitch5 * 2 + cwid, Inches(0.3))
para(tf, "evicted — no longer available", size=10.5, color=RED_D, bold=True, first=True, align=PP_ALIGN.CENTER)

# question -> confident partial answer
qy = Inches(3.62); qh = Inches(0.62)
sp, tf = box(s, ML, qy, Inches(4.7), qh, fill=BAND, line_c=HAIRLINE)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, "“Now list ALL cross-document inconsistencies.”", size=12.5, color=INK, italic=True, first=True)
arrow(s, ML + Inches(4.82), qy + Inches(0.31), ML + Inches(5.5), qy + Inches(0.31), color=SECOND, w=1.75)
sp, tf = box(s, ML + Inches(5.62), qy, CW - Inches(5.62), qh, fill=CARD, line_c=RED, line_w=1.5)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, [("Fluent, confident answer — built only on chunks 4–8. ", {"color": INK}),
          ("Nothing signals that chunks 1–3 are gone.", {"color": RED_D, "bold": True})],
     size=12.5, first=True)

# why this is its own problem
cy5b = Inches(4.62); ch5b = Inches(1.32)
card(s, ML, cy5b, Inches(5.9), ch5b, VIOLET, "Not the attention problem from earlier",
     ["Attention = how unevenly the model reads what is in front of it. Memory = whether earlier content is in front of it at all. Fixing one doesn’t fix the other."],
     body_size=12)
card(s, ML + Inches(6.15), cy5b, Inches(5.9), ch5b, RED_D, "Everything you accumulate decays too",
     ["Your findings list, coverage checklist and “do not skip” rules live in the same finite window — and one drifted or dead session loses all of it."],
     body_size=12)

strip(s, Inches(6.18), [("If state can’t live in the conversation, it must live outside it — in files. ", {"bold": True}),
                        ("That constraint drives the whole design of Approach 2.", {})])
footer(s)

# ================================================================ SLIDE 5
s = slide_new()
header(s, "Risk register", "Five failure modes that matter for this task")

rows = [
    ("Hallucination", "model", "Claims a change or omission that isn’t in either document", "“No evidence, no finding” — every finding must quote both texts"),
    ("Silent truncation", "platform", "Parts of the file never reach the model at all", "Controlled sections + coverage checklist, confirmed manually"),
    ("Missed material changes", "both", "Subtle edits in middle sections, tables, annexes", "Batching + global reconciliation + final sweep"),
    ("Run-to-run inconsistency", "both", "Same prompt → different findings each run", "Repeat runs; track stability; log prompts, versions, outputs"),
    ("Tables & charts", "platform", "Complex fiscal / DSA layouts parsed poorly", "Extract tables separately; scripted, deterministic comparison"),
]
ty = Inches(1.72)
col_x = [ML, ML + Inches(3.15), ML + Inches(7.05)]
col_w = [Inches(3.0), Inches(3.75), Inches(5.04)]
hdr = ["Failure mode", "What happens", "Control"]
for j, htxt in enumerate(hdr):
    sp, tf = box(s, col_x[j], ty, col_w[j], Inches(0.42), fill=BLUE_DD, line_c=None, radius=0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, htxt, size=12.5, color=WHITE, bold=True, first=True)
rh = Inches(0.78)
tag_color = {"model": VIOLET, "platform": YELLOW_D, "both": RED_D}
for i, (name, tag, what, ctrl) in enumerate(rows):
    y = ty + Inches(0.5) + rh * i + Inches(0.06) * i
    fill = CARD if i % 2 == 0 else BAND
    sp, tf = box(s, col_x[0], y, col_w[0], rh, fill=fill, line_c=HAIRLINE, line_w=0.75, radius=0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, [(name, {"bold": True, "color": INK, "size": 12.5}),
              ("   " + tag.upper(), {"color": tag_color[tag], "size": 9, "bold": True})],
         size=12.5, first=True)
    sp, tf = box(s, col_x[1], y, col_w[1], rh, fill=fill, line_c=HAIRLINE, line_w=0.75, radius=0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, what, size=11.5, color=SECOND, first=True, line=1.05)
    sp, tf = box(s, col_x[2], y, col_w[2], rh, fill=fill, line_c=HAIRLINE, line_w=0.75, radius=0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, ctrl, size=11.5, color=SECOND, first=True, line=1.05)

_, tf = tb(s, ML, Inches(6.6), CW, Inches(0.4))
para(tf, [("Source of the risk:  ", {"color": MUTED}),
          ("MODEL", {"color": VIOLET, "bold": True}), (" = inherent to LLMs   ", {"color": MUTED}),
          ("PLATFORM", {"color": YELLOW_D, "bold": True}), (" = Copilot product layer   ", {"color": MUTED}),
          ("BOTH", {"color": RED_D, "bold": True})], size=10.5, first=True)
footer(s)

# ================================================================ SLIDE 6
s = slide_new()
header(s, "Quality-control lens", "Treat it like model validation, not magic")

_, tf = tb(s, ML, Inches(1.62), CW, Inches(0.45))
para(tf, "Banks already run this playbook for credit-scoring models: repeat testing for precision and recall, then human sign-off.",
     size=13.5, color=SECOND, first=True)

py6 = Inches(2.3); ph6 = Inches(2.5); pw6 = Inches(5.9)
sp, tf = box(s, ML, py6, pw6, ph6, fill=CARD, line_c=YELLOW, line_w=1.5)
para(tf, [("False positive", {"bold": True, "color": INK, "size": 16}), ("   ·  Type I", {"color": YELLOW_D, "size": 12, "bold": True})], first=True, space_after=6)
para(tf, "AI flags a change that isn’t actually material", size=13.5, color=SECOND, space_after=10)
para(tf, [("Cost:  ", {"bold": True, "color": INK}), ("reviewer time to dismiss it", {})], size=13.5, color=SECOND, space_after=4)
para(tf, "Annoying — but cheap and visible", size=12.5, color=MUTED, italic=True)

sp, tf = box(s, ML + Inches(6.15), py6, pw6, ph6, fill=CARD, line_c=RED, line_w=1.5)
para(tf, [("False negative", {"bold": True, "color": INK, "size": 16}), ("   ·  Type II", {"color": RED_D, "size": 12, "bold": True})], first=True, space_after=6)
para(tf, "AI misses a material change entirely", size=13.5, color=SECOND, space_after=10)
para(tf, [("Cost:  ", {"bold": True, "color": INK}), ("flawed assessment, wrong policy advice, reputational damage", {})], size=13.5, color=SECOND, space_after=4)
para(tf, "Dangerous — silent and expensive", size=12.5, color=RED_D, italic=True)

strip(s, Inches(5.3), [("Design choice: ", {"bold": True}),
                       ("tune the AI first pass for high recall (accept extra false positives) → human expert review restores precision.", {})],
      h=Inches(0.85))
footer(s)

# ================================================================ SLIDE 7
s = slide_new()
header(s, "Approach 1", "One mega-prompt in Copilot Chat")

# left: the prompt
lw = Inches(4.55)
sp, tf = box(s, ML, Inches(1.7), lw, Inches(3.6), fill=RGBColor(0xF4, 0xF4, 0xF1), line_c=HAIRLINE)
para(tf, "THE PROMPT (condensed)", size=10.5, color=MUTED, bold=True, first=True, space_after=8)
para(tf, "“As an expert IMF economist, compare the Policy Note and the Staff Report…", size=13, color=INK, italic=True, space_after=6, line=1.2)
para(tf, "…identify ALL changes — deviations, additions, omissions, tone, strength, emphasis…", size=13, color=INK, italic=True, space_after=6, line=1.2)
para(tf, "…extract everything, then prove 100% coverage.”", size=13, color=INK, italic=True, space_after=10, line=1.2)
para(tf, "One message, both PDFs attached. Minutes to run.", size=11.5, color=MUTED)

# right: three numbered reasons
rx = ML + lw + Inches(0.35); rw = CW - lw - Inches(0.35)
_, tf = tb(s, rx, Inches(1.66), rw, Inches(0.35))
para(tf, "Why “prove 100% coverage” cannot be honored", size=14.5, color=INK, bold=True, first=True)
reasons = [
    ("1", "Output budget", "A reply has a hard length cap. The model silently compresses: 3 table rows “for illustration”, “similar revisions elsewhere…”"),
    ("2", "Retrieval is query-driven", "A diffuse mega-prompt makes a weak search query — entire tables never enter the model’s view at all"),
    ("3", "Instruction decay", "“Do not skip anything” fades as the context grows; hedges and shortcuts creep back in"),
]
ry = Inches(2.12)
for num, t, b in reasons:
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, rx, ry + Inches(0.12), Inches(0.42), Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE_DD; circ.line.fill.background(); circ.shadow.inherit = False
    ctf = circ.text_frame; ctf.word_wrap = False
    ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, num, size=15, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
    _, tf = tb(s, rx + Inches(0.62), ry, rw - Inches(0.62), Inches(1.05))
    para(tf, t, size=13.5, color=INK, bold=True, first=True, space_after=2)
    para(tf, b, size=12, color=SECOND, line=1.08)
    ry += Inches(1.1)

strip(s, Inches(5.65), [("Result: ", {"bold": True}),
                        ("a fluent, confident memo built on the most salient ~30% of the documents — a smart skimmer with no audit trail.", {})],
      fill=RED_D)
footer(s)

# ================================================================ SLIDE 8
s = slide_new()
header(s, "Approach 2", "A pipeline that forces coverage")

stages = [
    ("Inventory", "Ledger of every text unit: sections, tables, footnotes, annexes"),
    ("Chunk & batch", "Controlled sections sized to what the model reliably handles"),
    ("Pair", "Thematic crosswalk PN ↔ SR — topics, not heading order"),
    ("Compare", "Findings per batch, each backed by verbatim quotes from both docs"),
    ("Reconcile", "Document-wide matching: footnote ↔ main text ↔ annex ↔ table note"),
    ("Validate", "Scripted quote checks; manifest closed; limitations stated up front"),
]
cy8 = Inches(1.95); chev_h = Inches(0.85)
chev_w = Inches(2.28); overlap = Inches(0.31)
pitch = chev_w - overlap
for i, (t, sub) in enumerate(stages):
    x = ML + pitch * i
    sp = s.shapes.add_shape(MSO_SHAPE.CHEVRON, x, cy8, chev_w, chev_h)
    sp.fill.solid(); sp.fill.fore_color.rgb = RAMP[i]
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = Inches(0.14) if i == 0 else Inches(0.36)
    tf.margin_right = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tcol = INK if i < 2 else WHITE
    para(tf, f"{i+1}. {t}", size=11.5, color=tcol, bold=True, first=True, align=PP_ALIGN.LEFT)
    # description below, kept inside this stage's column
    _, dtf = tb(s, x + Inches(0.12), cy8 + chev_h + Inches(0.18), pitch - Inches(0.24), Inches(1.7))
    para(dtf, sub, size=10.5, color=SECOND, first=True, line=1.1)

# files-as-state visual
fy8 = Inches(4.55)
_, tf = tb(s, ML, fy8, Inches(4.2), Inches(0.35))
para(tf, "Between every stage: a file on disk", size=13, color=INK, bold=True, first=True)
file_labels = ["inventory.csv", "batches/*.md", "crosswalk.csv", "findings/*.md", "recon_report.md", "manifest: all DONE"]
fx = ML
for i, fl in enumerate(file_labels):
    sp, tf = box(s, fx, fy8 + Inches(0.45), Inches(1.78), Inches(0.5), fill=CARD, line_c=RAMP[i], line_w=1.25, radius=0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    para(tf, fl, size=10, color=SECOND, bold=True, first=True, align=PP_ALIGN.CENTER)
    if i > 0:
        arrow(s, fx - Inches(0.25), fy8 + Inches(0.7), fx - Inches(0.02), fy8 + Inches(0.7), color=MUTED, w=1.25)
    fx += Inches(2.06)

strip(s, Inches(5.95), [("Files carry the state, not the conversation ", {"bold": True}),
                        ("— each stage starts a fresh, small chat context; a crashed or drifting session costs one step, not the whole run.", {})])
footer(s)

# ================================================================ SLIDE 9
s = slide_new()
header(s, "Design principles", "Why the pipeline looks like this")

pr9 = [
    ("Files are the memory", "Every intermediate product is written to disk. New chat per stage — the files carry the state, not the conversation window.", BLUE),
    ("Thematic crosswalk, not structural alignment", "A Policy Note and a Staff Report are not two versions of one document. Align by topic; coverage asymmetries are findings, not noise.", AQUA_D),
    ("Determinism wherever possible", "Section splitting, number extraction, pairing math, quote verification run as scripts with checkable output — model judgment only where it’s needed.", VIOLET),
    ("Completeness by manifest, not by whim", "A CSV lists every topic pair; every row must reach DONE; a final sweep pass catches anything the crosswalk missed.", BLUE_DD),
]
gx = [ML, ML + Inches(6.15)]
gy = [Inches(1.8), Inches(4.15)]
gw, gh = Inches(5.9), Inches(2.15)
for i, (t, b, c) in enumerate(pr9):
    card(s, gx[i % 2], gy[i // 2], gw, gh, c, f"{i+1}.  {t}", [b], title_size=15, body_size=13)
footer(s)

# ================================================================ SLIDE 10
s = slide_new()
header(s, "Why both approaches exist", "Same engine underneath — different guarantees")

rows10 = [
    ("Process", "Retrieve → analyze → synthesize, hidden and opportunistic", "The same steps — explicit, forced, logged"),
    ("Coverage", "Unknown; whatever retrieval happened to surface", "Manifest: every section pair visited"),
    ("Evidence", "Fluent prose; quotes unverifiable", "Verbatim quotes, verified by script"),
    ("Reproducibility", "Varies run to run", "Logged prompts, inputs and outputs"),
    ("Effort", "Minutes", "Hours — and reusable for the next review"),
]
ty = Inches(1.72)
cx10 = [ML, ML + Inches(2.3), ML + Inches(7.2)]
cw10 = [Inches(2.15), Inches(4.75), Inches(4.89)]
heads = ["", "Copilot Chat (mega-prompt)", "Structured pipeline"]
for j in (1, 2):
    sp, tf = box(s, cx10[j], ty, cw10[j], Inches(0.45), fill=(SECOND if j == 1 else BLUE_DD), line_c=None, radius=0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, heads[j], size=12.5, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
rh10 = Inches(0.72)
for i, (dim, a, b) in enumerate(rows10):
    y = ty + Inches(0.55) + (rh10 + Inches(0.06)) * i
    fill = CARD if i % 2 == 0 else BAND
    sp, tf = box(s, cx10[0], y, cw10[0], rh10, fill=None, line_c=None)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, dim, size=12.5, color=INK, bold=True, first=True)
    sp, tf = box(s, cx10[1], y, cw10[1], rh10, fill=fill, line_c=HAIRLINE, line_w=0.75, radius=0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, a, size=11.5, color=SECOND, first=True, line=1.05)
    sp, tf = box(s, cx10[2], y, cw10[2], rh10, fill=fill, line_c=HAIRLINE, line_w=0.75, radius=0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, b, size=11.5, color=SECOND, first=True, line=1.05)

strip(s, Inches(6.35), [("The pipeline turns an opaque, convenience-oriented process into a controlled, auditable QC workflow.", {"bold": True})],
      h=Inches(0.55), fill=BLUE_DD)
footer(s)

# ================================================================ SLIDE 11
s = slide_new()
header(s, "Honest limits", "What the pipeline still doesn’t guarantee")

lim = [
    ("Coverage ≠ depth", "The manifest proves every pair was visited — not read carefully. An agent can skim, write thin findings, tick the row.", "The box says “done”, not “done well”."),
    ("Quote checks ≠ recall", "Scripts catch fabricated evidence. They cannot catch an omitted finding — the tone shift the model never noticed.", "There is no script for “what’s missing”."),
    ("Compliance is probabilistic", "Agents drift: skipped steps, buggy extraction scripts, optimistic “complete”. First runs hit snags.", "Expect run → fix → resume (cheap: stages are files)."),
]
cx11 = ML
for t, b, tail in lim:
    sp, tf = box(s, cx11, Inches(1.8), Inches(3.9), Inches(3.3), fill=CARD, line_c=HAIRLINE)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx11, Inches(1.8), Inches(3.9), Inches(0.09))
    bar.fill.solid(); bar.fill.fore_color.rgb = YELLOW; bar.line.fill.background(); bar.shadow.inherit = False
    para(tf, t, size=15, color=INK, bold=True, first=True, space_before=4, space_after=8)
    para(tf, b, size=12.5, color=SECOND, line=1.18, space_after=10)
    para(tf, tail, size=12, color=YELLOW_D, bold=True, line=1.1)
    cx11 += Inches(4.1)

strip(s, Inches(5.55), [("The final control is unchanged: an expert economist. ", {"bold": True}),
                        ("AI widens recall; the human supplies judgment and precision.", {})],
      fill=AQUA_D, h=Inches(0.85))
footer(s)

# ================================================================ SLIDE 12
s = slide_new()
header(s, "Field notes", "Practical lessons from the build")

fn = [
    ("Digital PDFs need no OCR", "Born-digital files carry a native text layer — parsing is fast and ~100% accurate. When visual/OCR checks were blocked, the limitation was declared up front, not papered over.", AQUA_D),
    ("Narrow beats broad", "One focused command per turn = a sharp retrieval query, a small output that fits the budget, and freshly re-anchored instructions.", BLUE),
    ("Know the billing unit", "Copilot charges per chat message (premium request), not per token — so the efficient pattern is few, well-loaded turns.", VIOLET),
    ("Leave a trail", "Save prompts, dates, model/tool versions, inputs, outputs. Findings that aren’t stable across runs get extra human review.", BLUE_DD),
]
gx = [ML, ML + Inches(6.15)]
gy = [Inches(1.8), Inches(4.15)]
for i, (t, b, c) in enumerate(fn):
    card(s, gx[i % 2], gy[i // 2], Inches(5.9), Inches(2.15), c, t, [b], title_size=15, body_size=13)
footer(s)

# ================================================================ SLIDE 13
s = slide_new()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.14))
band.fill.solid(); band.fill.fore_color.rgb = BLUE_DD; band.line.fill.background(); band.shadow.inherit = False
header(s, "Takeaways", "If you remember four things")

tk = [
    ("Chat AI is a smart skimmer", "Great for orientation and drafting — unsafe as the basis for “we found everything”."),
    ("Harnesses make thoroughness possible, not automatic", "Left to its own judgment an agent satisfices; structure has to force coverage."),
    ("Tune the first pass for recall", "A missed material change costs far more than a false alarm; human review restores precision."),
    ("Auditability is the product", "Ledger, manifest and logs turn an opaque answer into a review someone can check — and defend."),
]
ty13 = Inches(1.9)
for i, (t, b) in enumerate(tk):
    y = ty13 + Inches(1.18) * i
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, ML, y, Inches(0.55), Inches(0.55))
    circ.fill.solid(); circ.fill.fore_color.rgb = RAMP[min(i + 1, 5)]
    circ.line.fill.background(); circ.shadow.inherit = False
    ctf = circ.text_frame
    ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, str(i + 1), size=18, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
    _, tf = tb(s, ML + Inches(0.85), y - Inches(0.05), CW - Inches(0.85), Inches(1.1))
    para(tf, t, size=17, color=INK, bold=True, first=True, space_after=2)
    para(tf, b, size=13.5, color=SECOND)
footer(s)

# strip theme style refs (fillRef/effectRef) so no inherited shadows render
for slide in prs.slides:
    for shp in slide.shapes:
        el = shp._element
        style = el.find(qn('p:style'))
        if style is not None:
            el.remove(style)

# ---------------------------------------------------------------- save
out = "/tmp/claude-0/-home-user-documents/7ecfda3e-8533-53f0-a7ef-c90134788ccb/scratchpad/imf_ai_comparison_slides.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
