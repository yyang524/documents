#!/usr/bin/env python3
"""Create saudi_blanchard_leigh_imf.xlsx from the vintage/actuals dataset."""
import sys
sys.path.insert(0, '/tmp/claude-0/-home-user-documents/2056b728-c5cf-5ac0-a7f4-7459adce6669/scratchpad')
from build_workbook import V, A_LATEST, A_FIRST, YEARS  # noqa: E402

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties

OUT = "/home/user/documents/saudi_blanchard_leigh_imf.xlsx"

BLUE = "2A78D6"; VIOLET = "4A3AA7"; GRAY = "6B7280"; INKDARK = "1F2937"
HDR_FILL = PatternFill("solid", fgColor="EEF2F7")
BOLD = Font(bold=True, color=INKDARK)
H1 = Font(bold=True, size=14, color=INKDARK)
H2 = Font(bold=True, size=11, color=INKDARK)
NOTEFONT = Font(size=10, color="555555", italic=True)
THIN = Side(style="thin", color="D0D5DB")
BORDER = Border(bottom=THIN)

wb = Workbook()

# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 118
readme = [
 ("Saudi Arabia: Blanchard-Leigh (2013) exercise on IMF non-oil projections", H1),
 ("", None),
 ("Purpose", H2),
 ("Blanchard & Leigh (2013, AER P&P 'Growth Forecast Errors and Fiscal Multipliers') regress growth-forecast errors on the", None),
 ("fiscal consolidation planned at forecast time: FE(t) = alpha + beta * planned consolidation(t) + eps. beta < 0 indicates", None),
 ("fiscal multipliers larger than assumed by forecasters. This workbook adapts the design to Saudi Arabia's non-oil economy:", None),
 ("  - Outcome: real NON-OIL GDP growth, year t.", None),
 ("  - Fiscal variable: NON-OIL PRIMARY BALANCE (NOPB), percent of non-oil GDP.", None),
 ("  - Planned consolidation for t = NOPB(t) projection minus NOPB(t-1) estimate, both from the same pre-budget vintage.", None),
 ("", None),
 ("Projection vintages ('official projections')", H2),
 ("The Saudi authorities do not publish non-oil GDP growth or NOPB projections consistently, so per the task instruction the", None),
 ("IMF staff projections closest before each budget's approval are used. Saudi budgets for year t are approved in Nov/Dec of", None),
 ("t-1; the IMF Article IV staff report for Saudi Arabia is published Jul-Oct, i.e. 2-5 months before the budget. The vintage", None),
 ("for budget year t is therefore the Article IV report published in t-1 (e.g. budget year 2017 <- CR 16/326 of Oct 2016).", None),
 ("There was no 2020 Article IV consultation (COVID), so budget year 2021 has no usable IMF non-oil vintage.", None),
 ("The IMF World Economic Outlook database was not used because it does not carry non-oil GDP or NOPB series.", None),
 ("", None),
 ("Actual outcomes", H2),
 ("Two variants are provided:", None),
 ("  1) LATEST IMF historical data (headline, per task instruction): the most recent staff report containing year t", None),
 ("     (CR 25/223 of Aug 2025 for 2022-2024; CR 24/280 for 2020-2021; CR 23/323 for 2019; CR 19/290 for 2017-2018).", None),
 ("  2) FIRST-OUTTURN estimates: year t as first shown as actual/estimate in the next available report. Included because", None),
 ("     GASTAT's 2023-2025 national-accounts rebasing changed definitions: CR 25/223 measures non-oil GDP as non-oil", None),
 ("     activities + government activities + net taxes, lifting 2022 non-oil growth from 5.3% (CR 24/280) to 10.9% and", None),
 ("     shifting NOPB/non-oil GDP from about -32% to -25%. Latest-vintage actuals are therefore NOT on the same basis as", None),
 ("     the older projection vintages; first-outturn actuals are definitionally closer to what forecasters targeted.", None),
 ("2025 actuals: IMF historical values not yet published (2026 Article IV report due ~Aug 2026). PR 26/181 (3 Jun 2026)", None),
 ("reports total real GDP growth of 4.5% in 2025 with 'continued strength in non-oil activities'; row left open.", None),
 ("", None),
 ("Results (see Regression sheet)", H2),
 ("With n = 7 usable observations (2017-2020, 2022-2024) the Blanchard-Leigh beta is statistically indistinguishable from", None),
 ("zero in every specification (full sample, excl. 2020 COVID year, excl. 2020+2022 rebasing year; latest or first-outturn", None),
 ("actuals). beta ranges from -0.31 to +0.30 with p-values 0.40-0.89. The sample is far too short for inference - as", None),
 ("anticipated, the scatter charts (Chart sheet and chart.png) are the more honest summary. Points sit both above and below", None),
 ("zero at similar planned consolidations; the dominant errors are oil-cycle/COVID events (2020) and data rebasing (2022),", None),
 ("not the size of planned fiscal adjustment.", None),
 ("", None),
 ("Files", H2),
 ("data/imf_reports/  - archived source PDFs (fetched from imf.org / elibrary.imf.org by GitHub Actions workflow).", None),
 ("Data sheet        - full dataset with formulas; Sources sheet - per-figure citations incl. table and page.", None),
 ("", None),
 ("Prepared 14 Jul 2026. All projection/actual values transcribed from Table 1 (Selected Economic Indicators) of the", NOTEFONT),
 ("cited IMF country reports; transcription verified against rendered table images.", NOTEFONT),
]
for i, (txt, font) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=txt)
    if font: c.font = font

# ------------------------------------------------------------------ Data
ws = wb.create_sheet("Data")
headers = [
    ("A", "Budget year t", 11),
    ("B", "IMF vintage (pub. before t's budget)", 30),
    ("C", "Vintage pub.", 11),
    ("D", "Proj. non-oil real GDP growth t (%)", 13),
    ("E", "Proj. NOPB t (% non-oil GDP)", 13),
    ("F", "Est. NOPB t-1 (% non-oil GDP)", 13),
    ("G", "Planned consolidation ΔNOPB (pp)", 13),
    ("H", "Actual non-oil growth t, latest IMF (%)", 13),
    ("I", "Actual NOPB t, latest IMF (%)", 13),
    ("J", "Latest-actual source", 26),
    ("K", "Forecast error, latest (pp) = H - D", 13),
    ("L", "Actual non-oil growth t, first outturn (%)", 13),
    ("M", "Actual NOPB t, first outturn (%)", 13),
    ("N", "Actual NOPB t-1, same report (%)", 13),
    ("O", "Actual ΔNOPB, first outturn (pp)", 13),
    ("P", "First-outturn source", 26),
    ("Q", "Forecast error, first outturn (pp) = L - D", 13),
]
for col, title, width in headers:
    c = ws[f"{col}1"]
    c.value = title
    c.font = BOLD
    c.fill = HDR_FILL
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions[col].width = width
ws.freeze_panes = "B2"

r = 2
for t in YEARS:
    v, al, af = V[t], A_LATEST[t], A_FIRST[t]
    ws.cell(row=r, column=1, value=t)
    ws.cell(row=r, column=2, value=v["report"])
    ws.cell(row=r, column=3, value=v["pub"])
    ws.cell(row=r, column=4, value=v["g_proj"])
    ws.cell(row=r, column=5, value=v["nopb_t"])
    ws.cell(row=r, column=6, value=v["nopb_tm1"])
    ws.cell(row=r, column=7, value=f"=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),E{r}-F{r},\"\")")
    ws.cell(row=r, column=8, value=al["g"])
    ws.cell(row=r, column=9, value=al["nopb"])
    ws.cell(row=r, column=10, value=al["src"])
    ws.cell(row=r, column=11, value=f"=IF(AND(ISNUMBER(H{r}),ISNUMBER(D{r})),H{r}-D{r},\"\")")
    ws.cell(row=r, column=12, value=af["g"])
    ws.cell(row=r, column=13, value=af["nopb"])
    ws.cell(row=r, column=14, value=af["nopb_prev"])
    ws.cell(row=r, column=15, value=f"=IF(AND(ISNUMBER(M{r}),ISNUMBER(N{r})),M{r}-N{r},\"\")")
    ws.cell(row=r, column=16, value=af["src"])
    ws.cell(row=r, column=17, value=f"=IF(AND(ISNUMBER(L{r}),ISNUMBER(D{r})),L{r}-D{r},\"\")")
    for col in range(1, 18):
        cell = ws.cell(row=r, column=col)
        cell.border = BORDER
        if col in (4, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15, 17):
            cell.number_format = "0.0"
    r += 1
note = ws.cell(row=r + 1, column=1,
    value="Notes: 2021 - no IMF vintage (2020 Article IV skipped, COVID). 2025/2026 - IMF actuals not yet published. "
          "2022 latest actual & 2024 first outturn are on GASTAT's rebased national accounts (CR 25/223 basis), not the basis of older vintages. "
          "2019 'Actual NOPB t-1 same report' is n.a. because CR 21/149's table starts in 2019.")
note.font = NOTEFONT
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=17)

# ------------------------------------------------------------------ Regression
ws = wb.create_sheet("Regression")
ws.column_dimensions["A"].width = 46
for col in "BCDEFG":
    ws.column_dimensions[col].width = 13
ws["A1"] = "Blanchard-Leigh regression: FE(t) = alpha + beta * planned ΔNOPB(t)"
ws["A1"].font = H2
ws["A2"] = "Live formulas (SLOPE/INTERCEPT/LINEST over the complete-case block below). Delete a row from the block to test sensitivity."
ws["A2"].font = NOTEFONT

# complete-case block referencing Data sheet
ws["A4"] = "Complete observations"; ws["A4"].font = BOLD
hdr = ["Budget year", "Planned ΔNOPB (pp)", "FE latest (pp)", "FE first outturn (pp)"]
for j, h in enumerate(hdr):
    c = ws.cell(row=5, column=1 + j, value=h); c.font = BOLD; c.fill = HDR_FILL
    c.alignment = Alignment(wrap_text=True)
usable = [t for t in YEARS if V[t]["g_proj"] is not None and A_LATEST[t]["g"] is not None]
r0 = 6
for i, t in enumerate(usable):
    dr = 2 + YEARS.index(t)  # row on Data sheet
    ws.cell(row=r0 + i, column=1, value=t)
    ws.cell(row=r0 + i, column=2, value=f"=Data!G{dr}").number_format = "0.0"
    ws.cell(row=r0 + i, column=3, value=f"=Data!K{dr}").number_format = "0.0"
    ws.cell(row=r0 + i, column=4, value=f"=Data!Q{dr}").number_format = "0.0"
r1 = r0 + len(usable) - 1

stats_start = r1 + 2
ws.cell(row=stats_start, column=1, value="Statistic").font = BOLD
ws.cell(row=stats_start, column=2, value="vs latest actuals").font = BOLD
ws.cell(row=stats_start, column=3, value="vs first outturn").font = BOLD
X = f"B{r0}:B{r1}"
for j, ycol in enumerate(["C", "D"]):
    Y = f"{ycol}{r0}:{ycol}{r1}"
    col = 2 + j
    vals = [
        ("beta (slope)",        f"=SLOPE({Y},{X})"),
        ("s.e. of beta",        f"=INDEX(LINEST({Y},{X},TRUE,TRUE),2,1)"),
        ("t-stat of beta",      f"=SLOPE({Y},{X})/INDEX(LINEST({Y},{X},TRUE,TRUE),2,1)"),
        ("alpha (intercept)",   f"=INTERCEPT({Y},{X})"),
        ("R-squared",           f"=RSQ({Y},{X})"),
        ("n",                   f"=COUNT({Y})"),
    ]
    for i, (label, formula) in enumerate(vals):
        ws.cell(row=stats_start + 1 + i, column=1, value=label)
        c = ws.cell(row=stats_start + 1 + i, column=col, value=formula)
        c.number_format = "0.000"
pre = stats_start + 8
ws.cell(row=pre, column=1,
        value="Pre-computed (statsmodels OLS): latest actuals, full n=7: beta=-0.21 (se 1.11, p=0.86); "
              "excl. 2020: beta=-0.31 (p=0.75); excl. 2020&2022: beta=+0.06 (p=0.89).").font = NOTEFONT
ws.cell(row=pre + 1, column=1,
        value="First-outturn actuals, full n=7: beta=+0.30 (se 0.65, p=0.67); excl. 2020: +0.22 (p=0.51); "
              "excl. 2020&2022: +0.30 (p=0.40). No specification is significant; sample too short for inference.").font = NOTEFONT

# ------------------------------------------------------------------ Chart sheet (native scatter)
wsc = wb.create_sheet("Chart")
wsc["A1"] = "Non-oil growth forecast error vs planned change in non-oil primary balance (IMF pre-budget vintages)"
wsc["A1"].font = H2
chart = ScatterChart()
chart.title = "Saudi Arabia: Blanchard-Leigh scatter (2017-2024 budget years)"
chart.style = None
chart.x_axis.title = "Planned consolidation: projected ΔNOPB, pp of non-oil GDP"
chart.y_axis.title = "Non-oil growth forecast error, pp"
chart.height = 12; chart.width = 20
chart.legend.position = "b"
xref = Reference(ws, min_col=2, min_row=r0, max_row=r1)

s1 = Series(Reference(ws, min_col=3, min_row=r0, max_row=r1), xref, title="vs latest IMF actuals")
s1.marker = Marker(symbol="circle", size=8)
s1.marker.graphicalProperties = GraphicalProperties(solidFill=BLUE)
s1.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
s1.trendline = Trendline(trendlineType="linear")
chart.series.append(s1)

s2 = Series(Reference(ws, min_col=4, min_row=r0, max_row=r1), xref, title="vs first-outturn actuals")
s2.marker = Marker(symbol="diamond", size=8)
s2.marker.graphicalProperties = GraphicalProperties(solidFill=VIOLET)
s2.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
chart.series.append(s2)

wsc.add_chart(chart, "A3")
wsc["A29"] = "Interactive Excel chart: series 1 (blue circles) uses latest IMF actuals with a linear trendline; series 2 (violet diamonds) uses first-outturn actuals."
wsc["A29"].font = NOTEFONT
wsc["A30"] = "Static labeled version with year annotations: see chart.png in the repository."
wsc["A30"].font = NOTEFONT

# ------------------------------------------------------------------ Sources
ws = wb.create_sheet("Sources")
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 30
ws.column_dimensions["E"].width = 78
src_rows = [
    ("Role", "Document", "Published", "Values used", "URL / archive"),
    ("Vintage t=2017", "IMF CR 16/326, 2016 Article IV staff report", "Oct 2016", "Table 1 p.41: non-oil growth 2016-17; NOPB 2016-17",
     "https://www.imf.org/en/Publications/CR/Issues/2016/12/31/Saudi-Arabia-2016-Article-IV-Consultation-Press-Release-Staff-Report-and-Informational-Annex-44328 | data/imf_reports/imf_sau_artiv_2016_cr16326.pdf"),
    ("Vintage t=2018", "IMF CR 17/316, 2017 Article IV staff report", "Oct 2017", "Table 1 p.41: non-oil growth 2017-18; NOPB 2017-18",
     "https://www.imf.org/en/Publications/CR/Issues/2017/10/05/Saudi-Arabia-2017-Article-IV-Consultation-Press-Release-and-Staff-Report-45312 | data/imf_reports/imf_sau_artiv_2017_cr17316.pdf"),
    ("Vintage t=2019", "IMF CR 18/263, 2018 Article IV staff report", "Aug 2018", "Table 1 p.41: non-oil growth 2018-19; NOPB 2018-19",
     "https://www.imf.org/en/Publications/CR/Issues/2018/08/24/Saudi-Arabia-2018-Article-IV-Consultation-Press-Release-and-Staff-Report-46195 | data/imf_reports/imf_sau_artiv_2018_cr18263.pdf"),
    ("Vintage t=2020", "IMF CR 19/290, 2019 Article IV staff report", "Sep 2019", "Table 1 p.44: non-oil growth 2019-20; NOPB 2019-20; also latest actuals for 2017-18",
     "https://www.imf.org/en/Publications/CR/Issues/2019/09/09/Saudi-Arabia-2019-Article-IV-Consultation-Press-Release-and-Staff-Report-48659 | data/imf_reports/imf_sau_artiv_2019_cr19290.pdf"),
    ("(t=2021)", "No 2020 Article IV consultation (COVID-19); no IMF non-oil vintage near the Dec-2020 budget", "-", "-", "-"),
    ("Vintage t=2022", "IMF CR 21/149, 2021 Article IV staff report", "Jul 2021", "Table 1 p.40: non-oil growth 2021-22; NOPB 2021-22; first outturn for 2019-20",
     "https://www.imf.org/en/Publications/CR/Issues/2021/07/07/Saudi-Arabia-2021-Article-IV-Consultation-Press-Release-and-Staff-Report-461736 | data/imf_reports/imf_sau_artiv_2021_cr21149.pdf"),
    ("Vintage t=2023", "IMF CR 22/274, 2022 Article IV staff report", "Aug 2022", "Table 1 p.45: non-oil growth 2022-23; NOPB 2022-23; first outturn for 2021",
     "https://www.imf.org/en/publications/cr/issues/2022/08/11/saudi-arabia-2022-article-iv-consultation-press-release-and-staff-report-522189 | data/imf_reports/imf_sau_artiv_2022_cr22274.pdf"),
    ("Vintage t=2024", "IMF CR 23/323, 2023 Article IV staff report", "Sep 2023", "Table 1 p.49: non-oil growth 2023-24; NOPB 2023-24; latest actuals for 2019; first outturn for 2022",
     "https://www.imf.org/en/publications/cr/issues/2023/09/05/saudi-arabia-2023-article-iv-consultation-press-release-staff-report-and-informational-annex-538823 | data/imf_reports/imf_sau_artiv_2023_cr23323.pdf"),
    ("Vintage t=2025", "IMF CR 24/280, 2024 Article IV staff report", "Sep 2024", "Table 1 p.52: non-oil growth 2024-25; NOPB 2024-25; latest actuals for 2020-21; first outturn for 2023",
     "https://www.imf.org/en/Publications/CR/Issues/2024/09/03/Saudi-Arabia-2024-Article-IV-Consultation-Press-Release-Staff-Report-and-Informational-Annex-554530 | data/imf_reports/imf_sau_artiv_2024_cr24280.pdf"),
    ("Vintage t=2026", "IMF CR 25/223, 2025 Article IV staff report", "Aug 2025", "Table 1 p.43: non-oil growth 2025-26; NOPB 2025-26; latest actuals for 2022-24 (rebased basis)",
     "https://www.imf.org/en/publications/cr/issues/2025/08/02/saudi-arabia-2025-article-iv-consultation-press-release-and-staff-report-569252 | data/imf_reports/imf_sau_artiv_2025_cr25223.pdf"),
    ("2025 memo", "IMF Press Release 26/181, staff completes 2026 Article IV mission", "3 Jun 2026", "Real GDP +4.5% in 2025 (total; non-oil actual not yet published)",
     "https://www.imf.org/en/news/articles/2026/06/03/pr26181-saudi-arabia-imf-staff-completes-2026-article-iv-mission"),
    ("Method", "Blanchard, O. and D. Leigh (2013), 'Growth Forecast Errors and Fiscal Multipliers', AER 103(3) / IMF WP/13/1", "2013", "Regression design", "https://www.imf.org/external/pubs/ft/wp/2013/wp1301.pdf"),
]
for i, row in enumerate(src_rows, 1):
    for j, val in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=val)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = BOLD; c.fill = HDR_FILL

wb.save(OUT)
print("saved", OUT)
